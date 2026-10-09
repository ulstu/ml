from pathlib import Path
from shutil import copyfile
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HERE = Path(__file__).resolve().parent
RAW = HERE / "article_en_springer.docx"
REPAIRED = HERE / "article_en_springer_repaired.docx"
FINAL = HERE / "article_en_springer_final.docx"


def strip_broken_embedded_font_relationships(src: Path, dst: Path) -> None:
    with ZipFile(src, "r") as zin, ZipFile(dst, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == "word/_rels/fontTable.xml.rels":
                continue
            data = zin.read(item.filename)
            if item.filename == "word/fontTable.xml":
                text = data.decode("utf-8")
                for tag in [
                    "w:embedRegular",
                    "w:embedBold",
                    "w:embedItalic",
                    "w:embedBoldItalic",
                ]:
                    while f"<{tag}" in text:
                        start = text.index(f"<{tag}")
                        end = text.index("/>", start) + 2
                        text = text[:start] + text[end:]
                data = text.encode("utf-8")
            zout.writestr(item, data)


def set_cell_width(cell, cm):
    cell.width = Cm(cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm / 2.54 * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, val="nil"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "auto")


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def is_equation_table(table) -> bool:
    if len(table.columns) != 2 or len(table.rows) != 1:
        return False
    right_text = table.cell(0, 1).text.strip()
    return right_text.startswith("(") and right_text.endswith(")") and right_text[1:-1].isdigit()


def style_document(path: Path, out: Path) -> None:
    doc = Document(path)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name in {"Body Text", "bobytext", "First Paragraph", "Compact", "Normal"}:
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.08
        if paragraph.style.name.lower().startswith("heading"):
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            if paragraph.style.name.lower().startswith("heading"):
                run.font.size = Pt(12)
            elif paragraph.style.name in {"reference"}:
                run.font.size = Pt(9)
            else:
                run.font.size = Pt(10)

    for table in doc.tables:
        if is_equation_table(table):
            set_table_borders(table, "nil")
            table.autofit = False
            for row in table.rows:
                set_cell_width(row.cells[0], 13.2)
                set_cell_width(row.cells[1], 1.2)
                row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in row.cells[0].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for p in row.cells[1].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_cell_margins(row.cells[0], 10, 0, 10, 0)
                set_cell_margins(row.cells[1], 10, 0, 10, 0)
        else:
            table.autofit = True
            for row in table.rows:
                for cell in row.cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell, 70, 90, 70, 90)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        for run in paragraph.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(8.5)

    doc.save(out)


def main():
    strip_broken_embedded_font_relationships(RAW, REPAIRED)
    style_document(REPAIRED, FINAL)
    print(FINAL)


if __name__ == "__main__":
    main()
